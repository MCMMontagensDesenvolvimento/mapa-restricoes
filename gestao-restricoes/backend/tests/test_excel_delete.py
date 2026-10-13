import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table

from app import sync_service


class ExcelDeleteTests(unittest.TestCase):
    def make_workbook(self, ids):
        temp_root = Path(__file__).parent / ".tmp"
        temp_root.mkdir(exist_ok=True)
        folder = tempfile.TemporaryDirectory(dir=temp_root)
        path = Path(folder.name) / "restricoes.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = sync_service.SHEET
        for column, header in sync_service.HEADERS.items():
            sheet.cell(6, column).value = header
        for row, restriction_id in enumerate(ids, 7):
            sheet.cell(row, 1).value = restriction_id
            sheet.cell(row, 6).value = f"Restrição {restriction_id}"
            sheet.cell(row, 13).value = "Aberta"
            sheet.cell(row, 15).value = f'=IF(K{row}="","",TODAY()-K{row})'
            sheet.cell(row, 16).value = f'=IF(L{row}="","",L{row}-TODAY())'
        if ids:
            sheet.add_table(Table(displayName="Restricoes", ref=f"A6:Q{6 + len(ids)}"))
        workbook.save(path)
        workbook.close()
        return folder, path

    def test_missing_id_preserves_workbook(self):
        folder, path = self.make_workbook(["A", "B"])
        with folder, patch.object(sync_service, "source_path", return_value=path), patch.object(sync_service, "backup_path", return_value=Path(folder.name) / "backups"):
            with self.assertRaisesRegex(ValueError, "não foi encontrada"):
                sync_service.delete_excel_restriction("X")
            workbook = load_workbook(path, read_only=True)
            self.assertEqual([workbook[sync_service.SHEET].cell(row, 1).value for row in (7, 8)], ["A", "B"])
            workbook.close()

    def test_duplicate_id_is_rejected_without_deleting_rows(self):
        folder, path = self.make_workbook(["DUP", "DUP"])
        with folder, patch.object(sync_service, "source_path", return_value=path), patch.object(sync_service, "backup_path", return_value=Path(folder.name) / "backups"):
            with self.assertRaisesRegex(ValueError, "mais de uma vez"):
                sync_service.delete_excel_restriction("DUP")
            workbook = load_workbook(path, read_only=True)
            self.assertEqual(sum(workbook[sync_service.SHEET].cell(row, 1).value == "DUP" for row in (7, 8)), 2)
            workbook.close()

    def test_deletes_exact_row_and_adjusts_table_and_formulas(self):
        folder, path = self.make_workbook(["A", "B", "C"])
        with folder, patch.object(sync_service, "source_path", return_value=path), patch.object(sync_service, "backup_path", return_value=Path(folder.name) / "backups"):
            result = sync_service.delete_excel_restriction("B")
            self.assertEqual((result["before"], result["after"]), (3, 2))
            workbook = load_workbook(path, data_only=False)
            sheet = workbook[sync_service.SHEET]
            self.assertEqual([sheet.cell(row, 1).value for row in (7, 8)], ["A", "C"])
            self.assertEqual(sheet.tables["Restricoes"].ref, "A6:Q8")
            self.assertIn("K8", sheet.cell(8, 15).value)
            workbook.close()


if __name__ == "__main__":
    unittest.main()
