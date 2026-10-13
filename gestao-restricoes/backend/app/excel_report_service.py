import io
from collections import Counter
from datetime import date,datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment,Border,Font,PatternFill,Protection,Side
from openpyxl.utils import get_column_letter

from .report_service import DONE,criticality,day_delta,situation,sorted_rows,summary

ORANGE='FF5A14';DARK='173C33';CHARCOAL='4F4F4F';WHITE='FFFFFF';PALE='FFF3EC'
LOGO=Path(__file__).parent/'assets'/'logo-mcm-horizontal.png'
HEADERS=['ID','Área/CWA','CWP','Disciplina','Tarefa','Restrição','Categoria','Ação para Liberação','Responsável','Empresa','Data Levantamento','Prazo','Dias p/ Prazo','Status','Prioridade','Criticidade','Impacto / Observação']
WIDTHS=[12,20,16,20,36,42,20,48,24,22,18,16,17,20,16,16,42]

def _value(value):return value if value not in (None,'') else '—'
def _deadline_fill(r,ref):
 if r.status in DONE:return 'EDF4F1'
 d=day_delta(r,ref)
 return 'FDEAEA' if d<0 else 'FFF0DC' if d<=1 else 'FFF8D8' if d<=7 else WHITE
def _setup(ws,title,project,owner,ref,rows):
 ws.sheet_view.showGridLines=False
 if LOGO.exists():
  logo=Image(str(LOGO));logo.width=190;logo.height=46;ws.add_image(logo,'A1')
 ws.merge_cells('D1:Q1');ws['D1']=title;ws['D1'].font=Font(size=18,bold=True,color=DARK);ws['D1'].alignment=Alignment(horizontal='right',vertical='center')
 ws.merge_cells('D2:Q2');ws['D2']=f'Projeto: {project}   |   Responsável: {owner or "Todos"}   |   Referência: {ref.strftime("%d/%m/%Y")}   |   Emissão: {datetime.now().strftime("%d/%m/%Y %H:%M")}'
 ws['D2'].font=Font(size=9,color=CHARCOAL);ws['D2'].alignment=Alignment(horizontal='right')
 ws.row_dimensions[1].height=38;ws.row_dimensions[2].height=18
 stats=summary(rows,ref);cards=[('Total',stats['total']),('Vencidas',stats['overdue']),('Próximos 3 dias',sum(day_delta(r,ref)<=3 and r.status not in DONE for r in rows)),('Críticas',stats['critical']),('Alta prioridade',stats['high']),('Em tratamento',sum(r.status=='Em Tratamento' for r in rows))]
 for i,(label,value) in enumerate(cards):
  col=1+i*3;ws.merge_cells(start_row=4,start_column=col,end_row=4,end_column=col+2);ws.merge_cells(start_row=5,start_column=col,end_row=5,end_column=col+2)
  a=ws.cell(4,col,label);b=ws.cell(5,col,value)
  for cell in (a,b):cell.fill=PatternFill('solid',fgColor=PALE);cell.alignment=Alignment(horizontal='center');cell.border=Border(bottom=Side(style='thin',color=ORANGE))
  a.font=Font(size=9,bold=True,color=CHARCOAL);b.font=Font(size=16,bold=True,color=ORANGE)

def _table(ws,rows,ref,response_mode=False,start=8):
 headers=HEADERS+(['Nova Previsão','Comentário / Retorno do Responsável'] if response_mode else [])
 widths=WIDTHS+([18,42] if response_mode else [])
 for col,(header,width) in enumerate(zip(headers,widths),1):
  cell=ws.cell(start,col,header);cell.fill=PatternFill('solid',fgColor=DARK);cell.font=Font(bold=True,color=WHITE);cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True);cell.border=Border(bottom=Side(style='medium',color=ORANGE));ws.column_dimensions[get_column_letter(col)].width=width
 ws.row_dimensions[start].height=34
 for row_index,r in enumerate(sorted_rows(rows,ref),start+1):
  days=day_delta(r,ref)
  values=[r.external_id or r.id,_value(r.area),_value(r.cwp),_value(r.discipline),_value(r.task),_value(r.restriction),_value(r.category),_value(r.action),_value(r.owner),_value(r.company),r.raised_at,_value(r.due_at),'' if days==99999 else days,_value(r.status),_value(r.priority),criticality(r,ref),_value(r.impact)]
  if response_mode:values += ['', '']
  for col,value in enumerate(values,1):
   cell=ws.cell(row_index,col,value);cell.alignment=Alignment(vertical='top',wrap_text=True);cell.fill=PatternFill('solid',fgColor=_deadline_fill(r,ref));cell.border=Border(bottom=Side(style='hair',color='D9D9D9'))
   if col in (11,12,18):cell.number_format='dd/mm/yyyy'
   if col==15 and r.priority=='Alta':cell.fill=PatternFill('solid',fgColor='F4CCCC');cell.font=Font(bold=True,color='9C0006')
   if col==16 and criticality(r,ref)=='Crítica':cell.fill=PatternFill('solid',fgColor='FF5A14');cell.font=Font(bold=True,color=WHITE)
   if response_mode and col>len(HEADERS):cell.fill=PatternFill('solid',fgColor='FFF2CC');cell.protection=Protection(locked=False)
  ws.row_dimensions[row_index].height=42
 end=max(start+1,start+len(rows));ws.auto_filter.ref=f'A{start}:{get_column_letter(len(headers))}{end}';ws.freeze_panes=f'A{start+1}'
 ws.page_setup.orientation='landscape';ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0;ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.print_title_rows=f'{start}:{start}'
 if response_mode:ws.protection.sheet=True;ws.protection.autoFilter=True;ws.protection.sort=True

def _summary_sheet(wb,rows,ref,project):
 ws=wb.active;ws.title='Resumo';_setup(ws,'LEAN CONTROL — RELATÓRIO GERAL',project,None,ref,rows)
 ws['A8']='Indicador';ws['B8']='Quantidade'
 stats=summary(rows,ref);items=[('Total de restrições',len(rows)),('Pendências',sum(r.status not in DONE for r in rows)),('Vencidas',stats['overdue']),('Críticas',stats['critical']),('Alta prioridade',stats['high']),('Próximos 7 dias',stats['soon'])]
 for c in ws[8][:2]:c.fill=PatternFill('solid',fgColor=DARK);c.font=Font(bold=True,color=WHITE)
 for i,item in enumerate(items,9):ws.cell(i,1,item[0]);ws.cell(i,2,item[1])
 ws.column_dimensions['A'].width=34;ws.column_dimensions['B'].width=18
def _breakdown(wb,title,items):
 ws=wb.create_sheet(title);ws.sheet_view.showGridLines=False;ws.append([title,'Quantidade'])
 for c in ws[1]:c.fill=PatternFill('solid',fgColor=DARK);c.font=Font(bold=True,color=WHITE)
 for key,count in sorted(Counter(_value(x) for x in items).items(),key=lambda x:(-x[1],x[0])):ws.append([key,count])
 ws.column_dimensions['A'].width=42;ws.column_dimensions['B'].width=16;ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions

def xlsx_bytes(rows,ref:date,project,owner=None,response_mode=False,general=False):
 rows=sorted_rows(rows,ref);wb=Workbook()
 if general:
  _summary_sheet(wb,rows,ref,project);ws=wb.create_sheet('Restrições');_setup(ws,'LEAN CONTROL — RELATÓRIO GERAL DE RESTRIÇÕES',project,None,ref,rows);_table(ws,rows,ref,response_mode)
  _breakdown(wb,'Por Responsável',[r.owner for r in rows]);_breakdown(wb,'Por Disciplina',[r.discipline for r in rows]);_breakdown(wb,'Por Área',[r.area for r in rows])
 else:
  ws=wb.active;ws.title='Restrições';_setup(ws,'LEAN CONTROL — RELATÓRIO DE RESTRIÇÕES',project,owner,ref,rows);_table(ws,rows,ref,response_mode)
 wb.calculation.fullCalcOnLoad=True;wb.calculation.forceFullCalc=True;buf=io.BytesIO();wb.save(buf);return buf.getvalue()
