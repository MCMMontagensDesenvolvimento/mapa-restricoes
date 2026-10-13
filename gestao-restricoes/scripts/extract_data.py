import json
import uuid
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "Base de dados (Planilha)" / "Panorama_Semanal_Restricoes_Lean.xlsx"
OUTPUT = Path(__file__).resolve().parents[1] / "public" / "restrictions.json"


def iso(value):
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    return value


book = load_workbook(SOURCE, data_only=True, read_only=True)
sheet = book[book.sheetnames[0]]
headers = [str(cell.value or "").strip() for cell in sheet[6]]
keys = [
    "id", "area", "cwp", "discipline", "task", "restriction", "category",
    "action", "owner", "company", "raisedAt", "dueAt", "status", "priority",
    "aging", "daysToDue", "impact",
]

records = []
for row_number, row in enumerate(sheet.iter_rows(min_row=7, values_only=True), start=7):
    values = [iso(value) for value in row[:17]]
    if not any(value not in (None, "") for value in values):
        continue
    natural = "|".join(str(value or "") for value in values[1:14])
    values[0] = str(values[0] or uuid.uuid5(uuid.NAMESPACE_URL, f"lean:{natural}"))
    record = dict(zip(keys, values))
    record["sourceRow"] = row_number
    records.append(record)

lists = book[book.sheetnames[1]]
masters = {
    "categories": [row[0] for row in lists.iter_rows(min_row=2, values_only=True) if row[0]],
    "statuses": [row[1] for row in lists.iter_rows(min_row=2, values_only=True) if row[1]],
    "priorities": [row[2] for row in lists.iter_rows(min_row=2, values_only=True) if row[2]],
}

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
OUTPUT.write_text(json.dumps({"source": SOURCE.name, "headers": headers, "masters": masters, "records": records}, ensure_ascii=False), encoding="utf-8")
print(f"Exported {len(records)} records to {OUTPUT}")
