import asyncio, hashlib, json, os, shutil, sqlite3, threading, unicodedata, uuid
from copy import copy
from datetime import date, datetime, timezone
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.utils.cell import get_column_letter, range_boundaries
from sqlalchemy import delete, func, select, update
from sqlalchemy.ext.asyncio import AsyncSession
from .config import get_settings
from .models import AuditEvent, IdMapping, Restriction, SyncConflict, SyncQueue, SyncRun

SHEET="Panorama Restrições"
FIELDS=["id","area","cwp","discipline","task","restriction","category","action","owner","company","raised_at","due_at","status","priority","aging","days_to_due","impact"]
COMPARE=["area","cwp","discipline","task","restriction","category","action","owner","company","raised_at","due_at","status","priority","impact"]
COLUMNS={name:index+1 for index,name in enumerate(FIELDS)}
HEADERS={1:"ID Exclusiva",2:"Área / CWA",3:"CWP",4:"Disciplina",5:"Nome da Tarefa",6:"Restrição",7:"Categoria",8:"Ação para Liberação",9:"Responsável",10:"Empresa",11:"Data Levantamento",12:"Prazo",13:"Status",14:"Prioridade",15:"Aging (dias)",16:"Dias p/ Prazo",17:"Impacto / Observação"}
_lock=threading.Lock()

def source_path():
 s=get_settings();return Path(s.local_excel_path) if s.local_excel_path else Path(__file__).resolve().parents[3]/"Base de dados (Planilha)"/"Panorama_Semanal_Restricoes_Lean.xlsx"
def backup_path():
 configured=get_settings().local_backup_path;return Path(configured) if configured else Path(__file__).resolve().parents[2]/"backups"
def cell(value):return value.date() if isinstance(value,datetime) else value
def clean_text(value):
 if value is None:return None
 return " ".join(str(value).replace("\u200b","").replace("\ufeff","").strip().split()) or None
def normalize_id(value):
 if value is None:return ""
 if isinstance(value,float) and value.is_integer():return str(int(value))
 return clean_text(value) or ""
def normalized(value):
 if isinstance(value,(date,datetime)):return cell(value).isoformat()
 text=clean_text(value)
 return unicodedata.normalize("NFKD",text or "").encode("ascii","ignore").decode().casefold()
def logical_hash(data):
 keys=("area","cwp","discipline","task","restriction","owner","raised_at")
 return hashlib.sha256("|".join(normalized(data.get(k)) for k in keys).encode()).hexdigest()
def record(values):
 data=dict(zip(FIELDS,[cell(v) for v in values[:17]]))
 if not data["id"]:
  natural="|".join(str(data.get(k) or "") for k in ("area","discipline","task","restriction","owner","raised_at"));data["id"]=str(uuid.uuid5(uuid.NAMESPACE_URL,"lean-control:"+natural))
 data["id"]=normalize_id(data["id"]);data.pop("aging",None);data.pop("days_to_due",None)
 for key in COMPARE:
  if key not in ("raised_at","due_at"):data[key]=clean_text(data.get(key))
 data["restriction"]=str(data.get("restriction") or "Restrição sem descrição");data["status"]=str(data.get("status") or "Aberta");data["priority"]=str(data.get("priority") or "Média");data["logical_hash"]=logical_hash(data);return data
def payload(r):return {**{k:getattr(r,k) for k in COMPARE},"id":r.external_id}
def json_default(value):return value.isoformat() if isinstance(value,(date,datetime)) else str(value)
def daily_backup(source,force=False):
 folder=backup_path();folder.mkdir(parents=True,exist_ok=True);stamp=datetime.now().strftime("%Y-%m-%d_%H%M%S" if force else "%Y-%m-%d");target=folder/f"Panorama_Restricoes_{stamp}.xlsx"
 if force or not target.exists():shutil.copy2(source,target)
 return target
def get_sheet(workbook):
 if SHEET not in workbook.sheetnames:raise ValueError(f"A aba obrigatória '{SHEET}' não foi encontrada.")
 sheet=workbook[SHEET]
 for column,expected in HEADERS.items():
  if str(sheet.cell(6,column).value or "").strip()!=expected:raise ValueError(f"Estrutura inesperada no Excel: coluna {column} deveria ser '{expected}'.")
 return sheet
def copy_row(sheet,source_row,target_row):
 sheet.row_dimensions[target_row].height=sheet.row_dimensions[source_row].height
 for column in range(1,sheet.max_column+1):
  source,target=sheet.cell(source_row,column),sheet.cell(target_row,column)
  if source.has_style:target._style=copy(source._style)
  target.font=copy(source.font);target.fill=copy(source.fill);target.border=copy(source.border);target.alignment=copy(source.alignment);target.protection=copy(source.protection);target.number_format=source.number_format
 for column in (15,16):
  formula=sheet.cell(source_row,column).value
  if isinstance(formula,str) and formula.startswith("="):target=sheet.cell(target_row,column);target.value=Translator(formula,origin=sheet.cell(source_row,column).coordinate).translate_formula(target.coordinate)
def find_row(sheet,restriction_id):
 matches=[row for row in range(7,sheet.max_row+1) if str(sheet.cell(row,1).value or "").strip()==restriction_id]
 if len(matches)>1:raise ValueError(f"A ID Exclusiva {restriction_id} aparece mais de uma vez no Excel.")
 if matches:return matches[0],False
 for row in range(7,sheet.max_row+1):
  if not any(sheet.cell(row,column).value not in (None,"") for column in [*range(1,15),17]):return row,True
 target=sheet.max_row+1;copy_row(sheet,target-1,target);return target,True
def write_excel(data):
 source=source_path()
 if not source.exists():raise FileNotFoundError(f"Não foi possível localizar o arquivo Excel configurado: {source}")
 with _lock:
  daily_backup(source);workbook=load_workbook(source,data_only=False,read_only=False);sheet=get_sheet(workbook);row,created=find_row(sheet,str(data["id"]))
  if created and row>7:copy_row(sheet,row-1,row)
  for field in ["id",*COMPARE]:
   value=data.get(field)
   if field in ("raised_at","due_at") and isinstance(value,str) and value:value=date.fromisoformat(value)
   sheet.cell(row,COLUMNS[field]).value=value
  if not str(sheet.cell(row,15).value or "").startswith("="):sheet.cell(row,15).value=f'=IF(K{row}="","",IF(M{row}="Concluída","",TODAY()-K{row}))'
  if not str(sheet.cell(row,16).value or "").startswith("="):sheet.cell(row,16).value=f'=IF(L{row}="","",L{row}-TODAY())'
  temporary=source.with_name(f".{source.stem}.{uuid.uuid4().hex}.tmp.xlsx")
  try:
   workbook.save(temporary);check_book=load_workbook(temporary,data_only=False,read_only=True);check_sheet=get_sheet(check_book);check_row=next((i for i in range(7,check_sheet.max_row+1) if str(check_sheet.cell(i,1).value or "").strip()==str(data["id"])),None)
   if check_row is None:raise IOError("O arquivo foi salvo, mas a ID não pôde ser confirmada.")
   for field in ["id",*COMPARE]:
    actual=cell(check_sheet.cell(check_row,COLUMNS[field]).value)
    if str(actual or "")!=str(data.get(field) or ""):raise IOError(f"Falha na conferência do campo {field} no Excel.")
   check_book.close();os.replace(temporary,source)
  finally:
   workbook.close()
   if temporary.exists():temporary.unlink()
  return {"row":row,"created":created,"path":str(source)}
def delete_excel_restriction(restriction_id):
 source=source_path();restriction_id=normalize_id(restriction_id)
 if not source.exists():raise FileNotFoundError(f"Não foi possível localizar o arquivo Excel configurado: {source}")
 with _lock:
  workbook=load_workbook(source,data_only=False,read_only=False);sheet=get_sheet(workbook)
  matches=[row for row in range(7,sheet.max_row+1) if normalize_id(sheet.cell(row,1).value)==restriction_id]
  if not matches:workbook.close();raise ValueError(f"A ID Exclusiva {restriction_id} não foi encontrada no Excel. Nenhum dado foi excluído.")
  if len(matches)>1:workbook.close();raise ValueError(f"A ID Exclusiva {restriction_id} aparece mais de uma vez no Excel. Nenhum dado foi excluído.")
  row=matches[0];before=sum(1 for current in range(7,sheet.max_row+1) if normalize_id(sheet.cell(current,1).value));backup=daily_backup(source,True)
  table_refs={table.name:table.ref for table in sheet.tables.values()}
  sheet.delete_rows(row,1)
  for table in sheet.tables.values():
   min_col,min_row,max_col,max_row=range_boundaries(table_refs[table.name])
   if min_row<=row<=max_row:table.ref=f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row-1}"
  for current in range(max(7,row),sheet.max_row+1):
   if normalize_id(sheet.cell(current,1).value):
    sheet.cell(current,15).value=f'=IF(K{current}="","",IF(M{current}="Concluída","",TODAY()-K{current}))'
    sheet.cell(current,16).value=f'=IF(L{current}="","",L{current}-TODAY())'
  temporary=source.with_name(f".{source.stem}.{uuid.uuid4().hex}.tmp.xlsx")
  try:
   workbook.save(temporary);check_book=load_workbook(temporary,data_only=False,read_only=False);check_sheet=get_sheet(check_book)
   remaining=[current for current in range(7,check_sheet.max_row+1) if normalize_id(check_sheet.cell(current,1).value)==restriction_id]
   after=sum(1 for current in range(7,check_sheet.max_row+1) if normalize_id(check_sheet.cell(current,1).value))
   if remaining or after!=before-1:raise IOError("A exclusão não pôde ser confirmada no Excel. Nenhum dado foi removido do banco.")
   for table in check_sheet.tables.values():
    if table.name in table_refs:
     _,minimum,_,maximum=range_boundaries(table.ref)
     if minimum<=row and maximum>=row and table.ref==table_refs[table.name]:raise IOError("A faixa da tabela do Excel não foi ajustada após a exclusão.")
   check_book.close();os.replace(temporary,source)
  finally:
   workbook.close()
   if temporary.exists():temporary.unlink()
  return {"row":row,"path":str(source),"backup":str(backup),"before":before,"after":after}
def backfill_ids():
 source=source_path();workbook=load_workbook(source,data_only=False,read_only=False);sheet=get_sheet(workbook);changed=0
 for row in range(7,sheet.max_row+1):
  values=[sheet.cell(row,column).value for column in range(1,18)]
  if not any(v not in (None,"") for v in values[:14]+values[16:]):continue
  if not values[0]:sheet.cell(row,1).value=record(values)["id"];changed+=1
 if changed:
  daily_backup(source,True);temporary=source.with_name(f".{source.stem}.{uuid.uuid4().hex}.tmp.xlsx")
  try:workbook.save(temporary);os.replace(temporary,source)
  finally:
   if temporary.exists():temporary.unlink()
 workbook.close();return changed

async def save_restriction(db:AsyncSession,data:dict,user,current:Restriction|None=None):
 created=current is None;data.pop("rowVersion",None)
 if created:
  existing_ids=(await db.scalars(select(Restriction.external_id))).all();numeric=[int(x) for x in existing_ids if str(x or "").isdigit()];external_id=str(max(numeric,default=0)+1);data["logical_hash"]=logical_hash(data);current=Restriction(external_id=external_id,**data,completed_at=datetime.now(timezone.utc) if data.get("status")=="Concluída" else None,updated_by=user.email,update_origin="system");db.add(current);await db.flush();changes=[("restriction","",current.restriction)];event_name="created"
 else:
  changes=[]
  for field,value in data.items():
   old=getattr(current,field)
   if str(old or "")!=str(value or ""):changes.append((field,old,value));setattr(current,field,value)
  if current.status=="Concluída" and not current.completed_at:
   current.completed_at=datetime.now(timezone.utc);changes.append(("completed_at","",current.completed_at))
  elif current.status!="Concluída" and current.completed_at:
   old_completed=current.completed_at;current.completed_at=None;changes.append(("completed_at",old_completed,""))
  current.logical_hash=logical_hash({k:getattr(current,k) for k in COMPARE});current.row_version+=1;current.updated_by=user.email;current.update_origin="system";event_name="updated"
 event_ids=[]
 for field,old,new in changes:
  event=AuditEvent(restriction_id=current.id,restriction_external_id=current.external_id,event=event_name,field=field,old_value=str(old or ""),new_value=str(new or ""),actor_name=user.name,actor_email=user.email,origin="system",excel_result="pending");db.add(event);await db.flush();event_ids.append(event.id)
 queue=SyncQueue(restriction_id=current.id,operation="upsert",payload=json.dumps(payload(current),default=json_default,ensure_ascii=False),actor_email=user.email);db.add(queue);await db.commit();await db.refresh(current)
 try:
  result=await asyncio.to_thread(write_excel,json.loads(queue.payload));queue.status="synced";queue.attempts+=1;queue.synced_at=datetime.now(timezone.utc)
  if event_ids:await db.execute(update(AuditEvent).where(AuditEvent.id.in_(event_ids)).values(excel_result="synced"))
  await db.commit();return current,"synced",result
 except Exception as exc:
  queue.attempts+=1;queue.last_error=str(exc);await db.commit();return current,"pending",{"error":str(exc)}
async def retry_pending(db):
 pending=(await db.scalars(select(SyncQueue).where(SyncQueue.status=="pending").order_by(SyncQueue.created_at))).all();synced=0
 for item in pending:
  try:
   await asyncio.to_thread(write_excel,json.loads(item.payload));item.status="synced";item.synced_at=datetime.now(timezone.utc);item.last_error=None;synced+=1
   await db.execute(update(AuditEvent).where(AuditEvent.restriction_id==item.restriction_id,AuditEvent.excel_result=="pending").values(excel_result="synced"))
  except Exception as exc:item.last_error=str(exc)
  item.attempts+=1
 pending_ids=select(SyncQueue.restriction_id).where(SyncQueue.status=="pending")
 await db.execute(update(AuditEvent).where(AuditEvent.excel_result=="pending",AuditEvent.restriction_id.not_in(pending_ids)).values(excel_result="synced"))
 await db.commit();remaining=await db.scalar(select(func.count()).select_from(SyncQueue).where(SyncQueue.status=="pending"));return {"retried":len(pending),"synced":synced,"pending":remaining or 0}
async def synchronize_local(db,user_email):
 source=source_path()
 if not source.exists():raise FileNotFoundError(f"Não foi possível localizar o arquivo Excel configurado: {source}")
 await retry_pending(db);backup=daily_backup(source,True);book=load_workbook(source,data_only=True,read_only=True);sheet=get_sheet(book);incoming=[record(row) for row in sheet.iter_rows(min_row=7,values_only=True) if any(v not in (None,"") for v in list(row[:14])+list(row[16:17]))];book.close();run=SyncRun(provider="excel-local",status="started",actor_email=user_email);db.add(run);await db.flush();counts={"analyzed":len(incoming),"new":0,"updated":0,"unchanged":0,"conflicts":0,"renamed":0}
 for data in incoming:
  external_id=data.pop("id");signature=data.pop("logical_hash");current=await db.scalar(select(Restriction).where(Restriction.external_id==external_id))
  if not current:
   candidates=(await db.scalars(select(Restriction).where(Restriction.logical_hash==signature))).all()
   if len(candidates)==1:
    current=candidates[0];db.add(IdMapping(old_internal_id=current.id,old_external_id=current.external_id,new_internal_id=current.id,new_external_id=external_id,confidence="high"));current.external_id=external_id;counts["renamed"]+=1
   elif candidates:
    db.add(SyncConflict(restriction_id=external_id,field="external_id",database_value=",".join(x.external_id for x in candidates),external_value=external_id));counts["conflicts"]+=1;continue
   else:db.add(Restriction(external_id=external_id,logical_hash=signature,**data,updated_by=user_email,update_origin="excel"));counts["new"]+=1;continue
  changed=[field for field in COMPARE if str(getattr(current,field) or "")!=str(data.get(field) or "")]
  if not changed:current.logical_hash=signature;counts["unchanged"]+=1;continue
  pending=await db.scalar(select(SyncQueue).where(SyncQueue.restriction_id==current.id,SyncQueue.status=="pending"))
  if pending:
   for field in changed:db.add(SyncConflict(restriction_id=current.id,field=field,database_value=str(getattr(current,field) or ""),external_value=str(data.get(field) or "")))
   counts["conflicts"]+=1;continue
  for field in changed:
   old=getattr(current,field);new=data.get(field);setattr(current,field,new);db.add(AuditEvent(restriction_id=current.id,event="synchronized",field=field,old_value=str(old or ""),new_value=str(new or ""),actor_email=user_email,origin="excel",excel_result="not_applicable"))
  current.logical_hash=signature;current.row_version+=1;current.updated_by=user_email;current.update_origin="excel";counts["updated"]+=1
 run.status="completed";run.finished_at=datetime.now(timezone.utc);run.summary=json.dumps(counts,ensure_ascii=False);await db.commit();pending_count=await db.scalar(select(func.count()).select_from(SyncQueue).where(SyncQueue.status=="pending"));return {**counts,"pending":pending_count or 0,"status":"completed","lastSync":run.finished_at.isoformat(),"backup":str(backup.name)}

def database_path():
 url=get_settings().database_url
 if not url.startswith("sqlite"):raise ValueError("O backup automático desta rotina exige SQLite local.")
 return Path(url.split("///",1)[1]).resolve()
def backup_database(rows):
 source=database_path();folder=backup_path();folder.mkdir(parents=True,exist_ok=True);stamp=datetime.now().strftime("%Y-%m-%d_%H%M%S");target=folder/f"lean_control_antes_releitura_{stamp}.db";export=folder/f"lean_control_antes_releitura_{stamp}.json"
 source_connection=sqlite3.connect(source);target_connection=sqlite3.connect(target)
 try:source_connection.backup(target_connection)
 finally:target_connection.close();source_connection.close()
 export.write_text(json.dumps(rows,default=json_default,ensure_ascii=False,indent=2),encoding="utf-8");return target,export
def read_authoritative_excel():
 source=source_path();book=load_workbook(source,data_only=True,read_only=True);sheet=get_sheet(book);items=[];lines={}
 for row_number,row in enumerate(sheet.iter_rows(min_row=7,max_col=17,values_only=True),7):
  if not any(row[index] not in (None,"") for index in (0,1,3,4,5,8)):continue
  external_id=normalize_id(row[0])
  if not external_id:book.close();raise ValueError(f"Linha válida sem ID Exclusiva: {row_number}")
  lines.setdefault(external_id.casefold(),[]).append(row_number);item=record(row);item["id"]=external_id;items.append(item)
 duplicates={items[0]["id"] if False else key:rows for key,rows in lines.items() if len(rows)>1}
 book.close()
 if duplicates:raise ValueError("IDs duplicadas encontradas no Excel: "+"; ".join(f"{key} (linhas {', '.join(map(str,rows))})" for key,rows in duplicates.items()))
 return items
async def resync_authoritative_excel(db:AsyncSession,user_email:str):
 incoming=await asyncio.to_thread(read_authoritative_excel);existing=(await db.scalars(select(Restriction))).all();before=len(existing)
 serialized=[{"internal_id":r.id,"external_id":r.external_id,**{k:getattr(r,k) for k in COMPARE}} for r in existing];database_backup,json_backup=await asyncio.to_thread(backup_database,serialized)
 by_external={normalize_id(r.external_id).casefold():r for r in existing};by_signature={}
 for r in existing:by_signature.setdefault(r.logical_hash or logical_hash({k:getattr(r,k) for k in COMPARE}),[]).append(r)
 chosen=[];mapped_old=set();mappings=[]
 for data in incoming:
  external_id=data.pop("id");signature=data.pop("logical_hash");exact=by_external.get(external_id.casefold());candidates=by_signature.get(signature,[]);selected=exact or (candidates[0] if len(candidates)==1 else None);internal_id=selected.id if selected else str(uuid.uuid4())
  for old in candidates:
   mapped_old.add(old.id);mappings.append((old,internal_id,external_id))
  if exact and exact.id not in mapped_old:mapped_old.add(exact.id);mappings.append((exact,internal_id,external_id))
  chosen.append(Restriction(id=internal_id,external_id=external_id,logical_hash=signature,**data,row_version=(selected.row_version+1 if selected else 1),updated_by=user_email,update_origin="excel"))
 await db.execute(delete(SyncQueue));await db.execute(delete(SyncConflict));await db.execute(delete(Restriction));await db.flush();db.expunge_all()
 for item in chosen:db.add(item)
 for old,new_internal,new_external in mappings:db.add(IdMapping(old_internal_id=old.id,old_external_id=old.external_id,new_internal_id=new_internal,new_external_id=new_external,confidence="high"))
 db.add(AuditEvent(restriction_id=None,event="authoritative_rebuild",field="restrictions",old_value=str(before),new_value=str(len(chosen)),actor_email=user_email,origin="excel",excel_result="not_applicable"));await db.commit()
 after=await db.scalar(select(func.count()).select_from(Restriction));distinct=await db.scalar(select(func.count(func.distinct(Restriction.external_id))))
 if after!=len(incoming) or distinct!=after:raise RuntimeError(f"Reconstrução não reconciliada: Excel={len(incoming)}, banco={after}, IDs únicas={distinct}")
 old_ids_removed=[r.external_id for r in existing if r.id not in {x.id for x in chosen}]
 return {"before":before,"excelValid":len(incoming),"after":after,"oldRecordsRemoved":before-after,"duplicatesEliminated":max(0,before-after),"oldIdsRemoved":old_ids_removed,"currentIds":[r.external_id for r in chosen],"mappedIds":len(mappings),"databaseBackup":str(database_backup),"jsonBackup":str(json_backup),"status":"completed"}
